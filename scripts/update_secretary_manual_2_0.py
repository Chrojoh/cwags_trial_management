from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "CWAGS_Trial_Secretary_Manual_1.9.docx"
OUTPUT = ROOT / "docs" / "CWAGS_Trial_Secretary_Manual_2.0.docx"


def insert_after(document, paragraph, text, style="Normal"):
    created = document.add_paragraph(text, style=style)
    paragraph._p.addnext(created._p)
    return created


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    paragraph.text = paragraph.text.replace(old, new)


def find_paragraph(document, exact_text):
    return next(p for p in document.paragraphs if p.text.strip() == exact_text)


def find_heading(document, exact_text, style_name):
    return next(
        p
        for p in document.paragraphs
        if p.text.strip() == exact_text and p.style.name == style_name
    )


def paragraph_index(document, paragraph):
    return next(
        index
        for index, candidate in enumerate(document.paragraphs)
        if candidate._p is paragraph._p
    )


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def remove_body_range(document, first_element, stop_element):
    body = document.element.body
    removing = False
    for element in list(body.iterchildren()):
        if element is first_element:
            removing = True
        if element is stop_element:
            break
        if removing:
            body.remove(element)


def prevent_table_row_splits(document):
    """Keep each reference/glossary row intact across Word page boundaries."""
    for table in document.tables:
        for row in table.rows:
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
                row_properties.append(OxmlElement("w:cantSplit"))


def main():
    document = Document(SOURCE)

    for paragraph in document.paragraphs:
        replace_text(paragraph, "Version 1.9 - Secretary Edition", "Version 2.0 - Secretary Edition")
        replace_text(paragraph, "Version 1.9 – Secretary Edition", "Version 2.0 – Secretary Edition")
        replace_text(paragraph, "5.7 Trial Application and Approval", "5.8 Trial Application and Approval")
        replace_text(paragraph, "5.8 Trial Collaborators and Staff Roles", "5.9 Trial Collaborators and Staff Roles")
        replace_text(paragraph, "5.9 Moving a Postponed Trial Day", "5.10 Moving a Postponed Trial Day")

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            replace_text(paragraph, "Version 1.9", "Version 2.0")

    contents_anchor = find_paragraph(document, "5.8 Trial Application and Approval")
    contents_anchor.text = "5.7 Premium List and Printable Entry Form"
    contents_anchor = insert_after(
        document,
        contents_anchor,
        "5.8 Trial Application and Approval",
        contents_anchor.style,
    )

    application_heading = find_heading(
        document,
        "5.8 Trial Application and Approval",
        "Heading 2",
    )
    premium_heading = application_heading.insert_paragraph_before(
        "5.7 Premium List and Printable Entry Form",
        style="Heading 2",
    )
    current = premium_heading
    current = insert_after(
        document,
        current,
        "Prepare the premium while the trial is still a draft. The trial application can be sent for C-WAGS approval before the public entry link is opened. The Premium List Builder combines trial setup information with editable host information, so dates, classes, rounds, judges and fees do not need to be typed twice.",
    )
    current = insert_after(
        document,
        current,
        "To start from a prior event, choose an earlier saved premium for the same club and select Use This Premium. This copies the editable blocks, including policies, contacts, veterinarian information, directions and nearby services. It does not replace the current trial schedule, dates, judges, fees, waiver or uploaded map. Review every copied section because contacts, policies and local services can change.",
    )
    current = insert_after(
        document,
        current,
        "Enter the exact street address used for GPS directions. Add arrival instructions when the correct entrance is not obvious. A secretary-reviewed JPG or PNG local map can also be uploaded. The map remains private and is embedded only in the generated premium.",
    )
    current = insert_after(
        document,
        current,
        "The premium schedule is organized by trial day and uses the judges and fees saved in trial setup. Optional sections remain blank when the host supplies no information. The waiver page is generated from the trial waiver used by the entry form, keeping the public and paper versions consistent.",
    )
    current = insert_after(
        document,
        current,
        "Use Preview Paper Entry Form to produce a printable option for competitors who do not enter online. The paper form separates each trial day, follows the standard C-WAGS class order and uses the premium color scheme. Paper entries received by the secretary are entered through Live Event so they follow the same capacity, waitlist and financial rules as online entries.",
    )
    current = insert_after(
        document,
        current,
        "Save Draft while information is being reviewed. Mark Premium Ready only after the schedule exists and the trial waiver is complete. Generate a fresh PDF after any trial setup or premium change so the distributed document matches the current trial.",
    )

    financial_heading = find_heading(document, "11. Financial Management", "Heading 1")
    first_financial_body = next(
        p for p in document.paragraphs[paragraph_index(document, financial_heading) + 1 :]
        if p.text.strip()
    )
    handler_heading = first_financial_body.insert_paragraph_before(
        "Handler accounts and multiple dogs",
        style="Heading 2",
    )
    current = handler_heading
    current = insert_after(
        document,
        current,
        "The Financial Summary is organized by handler rather than by individual dog. C-WAGS numbers use the year and owner portion, shown as YY-OOOO, to place the handler's dogs on one account. The final two digits continue to identify each dog. An entry waiting for a C-WAGS number joins the registered handler account when its verified contact information matches.",
    )
    current = insert_after(
        document,
        current,
        "Expand the dog list on the handler row to review each dog's status, accepted regular and FEO runs, waitlisted rounds, quoted fees, accepted fees, waivers and reduced-rate status. Waitlisted rounds show their possible promoted cost but do not increase the current accepted balance.",
    )
    current = insert_after(
        document,
        current,
        "The amount owed, payments, refunds and current balance belong to the handler account. Record a payment once against the handler total even when several dogs are entered. If one dog is later deleted while another dog remains, the handler payment history is retained on the surviving account.",
    )
    current = insert_after(
        document,
        current,
        "Fee waivers and judge or volunteer rates apply to all entries selected for that handler. Expand the dog list before changing a rate when the row shows Mixed. Confirm the intended treatment for every dog, then refresh and verify the handler total.",
    )

    closing_heading = find_heading(document, "13. Closing the Trial", "Heading 1")
    closing_body = next(
        p for p in document.paragraphs[paragraph_index(document, closing_heading) + 1 :]
        if p.text.strip()
    )
    package_heading = closing_body.insert_paragraph_before(
        "Post-trial package",
        style="Heading 2",
    )
    current = package_heading
    current = insert_after(
        document,
        current,
        "Use the Post-Trial Package after scoring and financial reconciliation are complete. The readiness check identifies submitted entries awaiting acceptance, pending registration numbers, placeholder judges, missing scores and handler accounts with outstanding balances.",
    )
    current = insert_after(
        document,
        current,
        "Resolve each readiness item before sending results. The package includes the official-format Excel results workbook and the supporting trial review documents. Open every generated file and confirm the trial name, dates, classes, judges, results and totals before submission.",
    )
    current = insert_after(
        document,
        current,
        "Generate Close to Titles and Award Confirmations before the current trial results are posted to the C-WAGS tracker. Those reports add this trial's qualifying results to the tracker totals already on file. After posting, the same results may be counted twice.",
    )

    # The source document already applies a page break through the Heading 1 style.
    # Remove the additional manual page-break paragraph before section 15 so Word
    # does not create an empty page between Troubleshooting and the workflow.
    workflow_heading = find_heading(document, "15. Complete Trial Workflow", "Heading 1")
    workflow_index = paragraph_index(document, workflow_heading)
    if workflow_index > 0:
        preceding = document.paragraphs[workflow_index - 1]
        if not preceding.text.strip() and 'w:type="page"' in preceding._p.xml:
            remove_paragraph(preceding)

    # Keep this screenshot title with the screenshot row instead of leaving the
    # title alone at the bottom of the preceding page.
    for table in document.tables:
        if "Quick status actions" not in " ".join(cell.text for row in table.rows for cell in row.cells):
            continue
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
        break

    # Remove two obsolete low-resolution screenshots that do not remain legible
    # in the Word/PDF render. The surrounding instructions retain the complete
    # operational guidance without covering text or extending beyond the page.
    for caption_text in (
        "Award Confirmations keeps title verification available without the hidden ribbon-estimator workflow.",
        "Break-Even Analysis: save the per-run C-WAGS assumption so the automated expense estimate remains available when the page is reopened.",
    ):
        caption = find_paragraph(document, caption_text)
        caption_index = paragraph_index(document, caption)
        if caption_index > 0:
            image_paragraph = document.paragraphs[caption_index - 1]
            if not image_paragraph.text.strip() and "w:drawing" in image_paragraph._p.xml:
                remove_paragraph(image_paragraph)
        remove_paragraph(caption)

    # The detailed chapters already follow the complete trial lifecycle. The
    # former section 15 repeated the entire process after Closing and
    # Troubleshooting, sending the reader backward to trial creation. Remove that
    # duplicate block and keep the reference material at the end.
    duplicate_workflow = find_heading(document, "15. Complete Trial Workflow", "Heading 1")
    page_reference = find_heading(document, "16. Page and Button Reference", "Heading 1")
    remove_body_range(document, duplicate_workflow._p, page_reference._p)

    contents_workflow = next(
        p
        for p in document.paragraphs
        if p.text.strip() == "15. Complete Trial Workflow" and p.style.name == "List Bullet"
    )
    remove_paragraph(contents_workflow)

    for paragraph in document.paragraphs:
        replace_text(paragraph, "16. Page and Button Reference", "15. Page and Button Reference")
        replace_text(paragraph, "17. Glossary", "16. Glossary")
        replace_text(paragraph, "18. Final Secretary Verification", "17. Final Secretary Verification")

    prevent_table_row_splits(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
