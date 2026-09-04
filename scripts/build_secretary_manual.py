from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterable, Tuple

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = ROOT / "docs" / "CWAGS_Trial_Secretary_Manual_Draft_0.3.docx"
SOURCE_IMAGES = ROOT / "docs" / "secretary-manual-assets"
ANNOTATED = SOURCE_IMAGES / "annotated"
OUTPUT_DOCX = ROOT / "docs" / "CWAGS_Trial_Secretary_Manual_1.3.docx"
REDUCED_RATE_SETTING = SOURCE_IMAGES / "raw" / "28-reduced-rate-setting.png"
REDUCED_RATE_CHECKBOX = SOURCE_IMAGES / "raw" / "29-jv-checkbox.png"
REDUCED_RATE_ANNOTATED = ANNOTATED / "28-reduced-rate-workflow-annotated.png"


IMAGE_MAP: Dict[str, Tuple[str, str, Tuple[str, ...]]] = {
    "Manual dashboard reference": (
        "01-secretary-dashboard.png",
        "Start from the secretary dashboard",
        ("Choose an assigned trial from the sidebar.", "Review entry and financial summary cards.", "Confirm the trial before opening operational pages."),
    ),
    "Create Trial form": (
        "02-create-trial-basic-info.png",
        "Create the trial",
        ("Enter the event and location details.", "Set the date range and default fees.", "Review the secretary contact fields."),
    ),
    "Trial Days setup": (
        "04-select-trial-days.png",
        "Select the actual trial days",
        ("Choose only dates on which classes will run.", "Review the selected-day list.", "Save and continue when the dates are correct."),
    ),
    "Judge assignments": (
        "06-rounds-judges-and-reset.png",
        "Configure rounds and judges",
        ("Choose the day and class.", "Assign the primary judge to each round.", "A reset round is shown as a half-round."),
    ),
    "Fee setup": (
        "05-choose-classes-and-levels.png",
        "Choose classes and class fees",
        ("Select the offered levels for each day.", "Confirm the regular and FEO fees.", "Set class capacity where applicable."),
    ),
    "Secretary Entries page": (
        "12-entry-management-populated.png",
        "Review incoming entries",
        ("Use the summary counts and search controls.", "Review fee, run count, and status.", "Open selections when a dog has waitlisted rounds."),
    ),
    "Live Event waitlist": (
        "14-running-order-capacity-waitlist.png",
        "Manage a round waitlist",
        ("Active dogs keep their running positions.", "Waitlisted dogs remain unassigned and uncharged.", "Promote only when capacity is available."),
    ),
    "Out-of-order promotion warning": (
        "13-waitlist-management.png",
        "Promote a lower waitlist position",
        ("The list remains ordered by entry time.", "Selecting a lower dog produces a warning.", "Proceed only when the priority exception is justified."),
    ),
    "Capacity increase during promotion": (
        "14-running-order-capacity-waitlist.png",
        "Increase capacity before promotion",
        ("Review active and waitlisted counts.", "Increase capacity deliberately.", "Capacity changes do not promote dogs automatically."),
    ),
    "Running order": (
        "24-running-order-with-results.png",
        "Review the running order",
        ("Positions belong only to active selections.", "Drag, arrows, or keyboard controls can reorder dogs.", "Saved results appear beside each dog."),
    ),
    "Score-sheet export preview": (
        "15-score-entry-blank.png",
        "Prepare score entry",
        ("Select the correct day, class, and round.", "Waitlisted dogs are not included.", "Only active competitors appear on the score sheet."),
    ),
    "Live Event overview": (
        "14-running-order-capacity-waitlist.png",
        "Operate the live event screen",
        ("Select a class card at the top.", "Use Running Order Setup for positions and status.", "Switch to Score Entry when the round is ready."),
    ),
    "Scent score entry": (
        "16-score-entry-completed.png",
        "Record scent results",
        ("Pass/Fail is the required result field.", "Scents, faults, and time are optional details.", "Save all scores and confirm the result display."),
    ),
    "Competitor financial summary": (
        "19-financial-summary.png",
        "Track fees and balances",
        ("Save the C-WAGS and reduced-rate settings first.", "Use J/V only for an approved judge or volunteer rate.", "Use Waive for a documented full waiver."),
    ),
    "Fee waiver or adjustment": (
        "19-financial-summary.png",
        "Document a fee adjustment",
        ("Review the original amount and resulting balance.", "Use waiver controls only with authorization.", "Confirm the change appears in the journal."),
    ),
    "Trial Journal list": (
        "17-activity-journal.png",
        "Use the activity journal",
        ("Filter by action type, person, dog, or date.", "Each event includes a timestamp and concise description.", "Open audit details when investigating a change."),
    ),
    "Journal detail": (
        "17-activity-journal.png",
        "Review before-and-after details",
        ("Confirm who made the change.", "Compare the previous and new values.", "Use the record to resolve later questions."),
    ),
    "Trial completion checklist": (
        "18-class-summary.png",
        "Validate results before closing",
        ("Compare entered runs with completed results.", "Resolve missing scores, absences, and corrections.", "Export final reports only after totals reconcile."),
    ),
    "Final reports and exports": (
        "18-class-summary.png",
        "Review and export final results",
        ("Confirm class totals and completion rates.", "Verify passes, fails, and absences.", "Retain the exported workbook with the trial records."),
    ),
    "Manual trial application reference": (
        "20-trial-application.png",
        "Review the trial application",
        ("Confirm derived host, dates, and location.", "Complete optional application-only fields.", "Download the completed PDF when ready."),
    ),
    "Manual time calculator reference": (
        "21-time-calculator.png",
        "Plan the event schedule",
        ("Enter minutes and seconds per run.", "Compare used time with the daily allotment.", "Recalculate after entry changes."),
    ),
    "Scent score-sheet export day selection": (
        "25-score-sheet-export-day-selection.png",
        "Export scent score sheets by trial day",
        ("Open the export from Running Order Setup.", "Choose the day being prepared.", "Review active competitors and positions before printing."),
    ),
    "Summary export overview": (
        "26-summary-export-overview.png",
        "Reconcile the summary before export",
        ("Start with Select All Classes.", "Compare runs with completed results.", "Export only after passes, fails, and absences reconcile."),
    ),
}


REMOVE_PLACEHOLDERS = {
    "Running-order export preview",
    "Rally score entry",
    "Obedience score entry",
    "Record payment",
}


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def redact_header(draw: ImageDraw.ImageDraw, width: int) -> None:
    # Opaque replacement: do not blur personal account information.
    draw.rectangle((width - 245, 0, width, 78), fill=(255, 255, 255))
    draw.ellipse((width - 228, 17, width - 190, 55), fill=(230, 86, 0))
    draw.text((width - 218, 25), "JS", fill="white", font=font(14, True))
    draw.text((width - 180, 18), "Jamie Secretary", fill=(30, 30, 30), font=font(15, True))
    draw.text((width - 180, 42), "Trial Secretary", fill=(90, 90, 90), font=font(12))


def annotate_image(source: Path, target: Path, title: str, notes: Iterable[str]) -> None:
    image = Image.open(source).convert("RGB")
    draw = ImageDraw.Draw(image)
    redact_header(draw, image.width)

    if source.name == "01-secretary-dashboard.png":
        draw.rectangle((300, 168, 665, 224), fill=(225, 78, 0))
        draw.text((304, 182), "Welcome back, Jamie!", fill="white", font=font(24, True))

    if source.name == "02-create-trial-basic-info.png":
        draw.rectangle((307, 328, 681, 370), fill=(255, 255, 255))
        draw.rectangle((709, 328, 1084, 370), fill=(255, 255, 255))
        draw.text((318, 340), "Jamie Secretary", fill=(40, 40, 40), font=font(16))
        draw.text((720, 340), "secretary@example.com", fill=(40, 40, 40), font=font(16))

    # Trial application also contains the signed-in account in the contact card.
    if source.name == "20-trial-application.png":
        draw.rectangle((790, 405, 1215, 480), fill=(255, 246, 195), outline=(230, 150, 50), width=2)
        draw.text((810, 418), "Jamie Secretary", fill=(30, 30, 30), font=font(16, True))
        draw.text((810, 445), "secretary@example.com", fill=(60, 60, 60), font=font(14))

    if source.name == "19-financial-summary.png":
        # Replace outdated helper copy so the screenshot agrees with the
        # documented distinction between a reduced rate and a full waiver.
        draw.rectangle((300, 562, 930, 604), fill=(255, 248, 205))
        draw.text(
            (306, 566),
            "Save the approved rate before selecting J/V on the payment row.",
            fill=(80, 80, 80),
            font=font(13),
        )
        draw.text(
            (306, 584),
            "Leave J/V clear for the full fee; use Waive for a documented full waiver.",
            fill=(80, 80, 80),
            font=font(13),
        )

    banner_h = 112
    canvas = Image.new("RGB", (image.width, image.height + banner_h), (255, 247, 214))
    canvas.paste(image, (0, banner_h))
    cdraw = ImageDraw.Draw(canvas)
    cdraw.rectangle((0, 0, image.width, banner_h), fill=(255, 243, 190))
    cdraw.rectangle((0, 0, 14, banner_h), fill=(230, 86, 0))
    cdraw.text((34, 14), title, fill=(67, 39, 18), font=font(23, True))
    x = 34
    for number, note in enumerate(notes, 1):
        y = 53 + (number - 1) * 20
        cdraw.ellipse((x, y - 2, x + 18, y + 16), fill=(230, 86, 0))
        cdraw.text((x + 5, y), str(number), fill="white", font=font(11, True))
        cdraw.text((x + 28, y - 1), note, fill=(55, 55, 55), font=font(14))
    cdraw.rectangle((0, banner_h, image.width - 1, image.height + banner_h - 1), outline=(220, 150, 65), width=3)
    target.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(target, quality=94)


def build_reduced_rate_workflow() -> None:
    """Create a focused two-panel image for the saved rate and J/V control."""
    top = Image.open(REDUCED_RATE_SETTING).convert("RGB").crop((410, 285, 1760, 620))
    bottom = Image.open(REDUCED_RATE_CHECKBOX).convert("RGB").crop((450, 760, 1760, 1235))
    width = 1350
    banner_h = 128
    gap = 42
    canvas = Image.new("RGB", (width, banner_h + top.height + gap + bottom.height), "white")
    draw = ImageDraw.Draw(canvas)
    draw.rectangle((0, 0, width, banner_h), fill=(255, 243, 190))
    draw.rectangle((0, 0, 14, banner_h), fill=(230, 86, 0))
    draw.text((34, 14), "Apply an approved reduced rate", fill=(67, 39, 18), font=font(23, True))
    instructions = (
        "Enter the approved rate per run and select Save.",
        "Find the eligible person and select the J/V checkbox.",
        "Verify that the total and current balance recalculate.",
    )
    for number, text in enumerate(instructions, 1):
        y = 52 + (number - 1) * 22
        draw.ellipse((34, y - 2, 52, y + 16), fill=(230, 86, 0))
        draw.text((39, y), str(number), fill="white", font=font(11, True))
        draw.text((62, y - 1), text, fill=(55, 55, 55), font=font(14))
    canvas.paste(top, (0, banner_h))
    canvas.paste(bottom, (20, banner_h + top.height + gap))
    draw = ImageDraw.Draw(canvas)
    # Focus boxes: rate + Save on the first panel; J/V checkbox on the second.
    draw.rounded_rectangle((1110, banner_h + 230, 1335, banner_h + 315), radius=10, outline=(230, 86, 0), width=7)
    bottom_y = banner_h + top.height + gap
    draw.rounded_rectangle((280, bottom_y + 215, 365, bottom_y + 305), radius=10, outline=(230, 86, 0), width=7)
    draw.text((22, banner_h + top.height + 8), "1  Save the rate first", fill=(230, 86, 0), font=font(18, True))
    draw.text((380, bottom_y + 240), "2  Tick J/V for the eligible row", fill=(230, 86, 0), font=font(18, True))
    REDUCED_RATE_ANNOTATED.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(REDUCED_RATE_ANNOTATED, quality=95)


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.text = ""


def placeholder_title(table) -> str | None:
    text = table.cell(0, 0).text.strip()
    if "[SCREENSHOT PLACEHOLDER]" not in text:
        return None
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return lines[-1] if lines else None


def build() -> None:
    ANNOTATED.mkdir(parents=True, exist_ok=True)
    build_reduced_rate_workflow()
    annotated_by_title: Dict[str, Path] = {}
    for title, (filename, callout_title, notes) in IMAGE_MAP.items():
        target = ANNOTATED / f"{Path(filename).stem}-annotated.png"
        annotate_image(SOURCE_IMAGES / filename, target, callout_title, notes)
        annotated_by_title[title] = target

    document = Document(SOURCE_DOCX)
    document.paragraphs[0].text = "C-WAGS Trial Secretary Manual"
    document.paragraphs[1].text = "Secretary Operations and Event-Day Reference Guide"
    document.paragraphs[2].text = "Version 1.3 - Secretary Edition"

    replacements = {
        "A Word table of contents can be generated in the final edition after screenshots and page numbers are stable.":
            "This manual is written solely for trial secretaries. It follows the secretary workflow from trial creation and approval through entry management, event-day operations, financial record keeping, reports, and closing.",
        "Financial views should distinguish:":
            "Financial views are record-keeping tools. The application does not process credit cards or transfer funds. Secretary records should distinguish:",
        "payment recorded": "offline payment status recorded",
    }
    for paragraph in document.paragraphs:
        if paragraph.text in replacements:
            paragraph.text = replacements[paragraph.text]

    publish_contents = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "6. Publishing and Managing Entries"
        and paragraph.style.name == "List Bullet"
    )
    publish_contents.insert_paragraph_before(
        "5.7 Trial Application and Approval", style="List Bullet"
    )

    # Appendix A was a capture checklist for the drafting phase, and Appendix B
    # was retained reference material. Neither belongs in the finished secretary
    # edition now that annotated screenshots are embedded beside the procedures.
    appendix = next(
        (
            p
            for p in document.paragraphs
            if p.text.strip().startswith("Appendix A.") and p.style.name == "Heading 1"
        ),
        None,
    )
    if appendix is not None:
        body = document._element.body
        removing = False
        for child in list(body):
            if child is appendix._element:
                removing = True
            if removing and child.tag.endswith("}p"):
                body.remove(child)
    for paragraph in list(document.paragraphs):
        if paragraph.text.strip().startswith(("Appendix A.", "Appendix B.")):
            paragraph._element.getparent().remove(paragraph._element)

    # Expand the financial section with the saved C-WAGS fee workflow. Insert
    # immediately before the Trial Journal heading so the source draft remains
    # reusable and the generated manual keeps its existing style hierarchy.
    journal_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "12. Trial Journal" and paragraph.style.name == "Heading 1"
    )
    financial_content = (
        ("Break-Even setup and the C-WAGS amount", "Heading 2"),
        (
            "The Break-Even page stores the trial's financial assumptions. The most important value for the C-WAGS amount is the current C-WAGS fee charged for each regular run. Enter the fee published for the event; do not hard-code an old rate or include a fee that applies to a different type of run.",
            "Normal",
        ),
        ("Open the trial and select Financial Summary.", "List Number"),
        ("Select the Break-Even Analysis tab.", "List Number"),
        ("In C-WAGS Fee (per run), enter the current amount charged for one regular run.", "List Number"),
        ("Review the other assumptions on the page, then select Save Configuration.", "List Number"),
        ("Wait for the save confirmation before leaving the page.", "List Number"),
        (
            "The configuration is saved for that trial. When the secretary returns later, the saved fee is loaded again and the Financial Summary calculates the C-WAGS amount from the current regular runs entered in the trial.",
            "Normal",
        ),
        ("How the calculation changes", "Heading 2"),
        ("Accepted regular runs count toward the C-WAGS amount.", "List Bullet"),
        ("A waitlisted round has a current fee and C-WAGS amount of $0 until it is promoted.", "List Bullet"),
        ("Promoting a waitlisted run makes it active and adds it to the calculation.", "List Bullet"),
        ("Removing, withdrawing, or scratching a run removes it when that status is not billable.", "List Bullet"),
        ("A waived regular entry still represents a run owed to C-WAGS; the club absorbs that cost.", "List Bullet"),
        ("FEO runs are kept separate from this regular-run C-WAGS calculation.", "List Bullet"),
        (
            "Example: if a trial has 120 active regular runs and the saved C-WAGS fee is $X per run, the displayed amount owed to C-WAGS is 120 x $X. Use the current published fee in place of $X.",
            "Normal",
        ),
        (
            "Recheck the total after late entries, waitlist promotions, withdrawals, scratches, or fee waivers. The application records financial information but does not collect or transfer payments.",
            "Normal",
        ),
        ("Volunteer and judge reduced entry rates", "Heading 2"),
        (
            "A judge or volunteer does not receive a reduced entry fee automatically. The secretary applies the host club's approved policy to the appropriate financial row. Configure the rate before recording payment so the displayed balance is the amount the person is expected to pay.",
            "Normal",
        ),
        ("Full fee: leave the J/V checkbox clear. The normal class entry fees remain in effect.", "List Bullet"),
        ("Reduced fee: save a Volunteer / Reduced Entry Rate and then select J/V on the approved person's row.", "List Bullet"),
        ("Full waiver: select Waive and enter a clear reason. A waiver is different from a reduced rate and is retained in the financial and journal records.", "List Bullet"),
        ("How to apply a reduced rate", "Heading 2"),
        ("Open Financial Summary and remain on Expenses & Payments.", "List Number"),
        ("Enter the approved amount in Volunteer / Reduced Entry Rate - Rate per run.", "List Number"),
        ("Select Save and wait for confirmation.", "List Number"),
        ("Find the judge or volunteer in the payment table and review the dogs and runs included in that row.", "List Number"),
        ("Select the J/V checkbox. The program recalculates the affected entry fees and current balance using the saved rate.", "List Number"),
        ("Verify the new total and balance before recording any offline payment.", "List Number"),
        (
            "Use the reduction only after the person and benefit have been confirmed under the host club's policy. Typical timing is after entries exist and the judge or volunteer assignment is confirmed, but before the payment is recorded. If eligibility changes, clear J/V to restore the normal entry rate. If the saved reduced rate is changed later, clear and reselect J/V for each previously marked row that must be recalculated.",
            "Normal",
        ),
        (
            "The visible reduced-rate field is the regular-run rate. Review any FEO selections in the same row before applying J/V and do not assume that the regular reduced rate is also the intended FEO rate. A $0 reduced rate should only be used deliberately; for a true full waiver, use Waive so the reason and waived amount are documented.",
            "Normal",
        ),
        (
            "Judge compensation entered under Trial Expenses is separate from a judge's reduced entry fee. Adding a judge expense does not mark the judge's own dog entries as reduced, and selecting J/V does not create or pay a judge expense.",
            "Normal",
        ),
    )
    for text, style in financial_content:
        journal_heading.insert_paragraph_before(text, style=style)

    reduced_picture = journal_heading.insert_paragraph_before()
    reduced_picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    reduced_picture.add_run().add_picture(str(REDUCED_RATE_ANNOTATED), width=Inches(5.9))
    caption = journal_heading.insert_paragraph_before(
        "Save the approved rate, then select J/V on the eligible competitor row and confirm the recalculated balance."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # The trial application is a pre-publication approval step, not a closing
    # report. Place it after configuration and directly before the publishing
    # workflow so a secretary sees the required sequence in context.
    publish_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip() == "6. Publishing and Managing Entries"
        and paragraph.style.name == "Heading 1"
    )
    application_content = (
        ("5.7 Trial Application and Approval", "Heading 2"),
        (
            "After the trial has been configured, generate the Trial Application while the trial is still in Draft. The application can be completed and submitted to C-WAGS before the public trial is published or opened for entries.",
            "Normal",
        ),
        ("Review the configured days, classes, rounds, judges, fees, capacities, host, dates, and location.", "List Number"),
        ("Open Trial Application and complete any application-only information.", "List Number"),
        ("Download the completed PDF and submit it to C-WAGS for approval.", "List Number"),
        ("Keep the trial in Draft while approval is pending.", "List Number"),
        ("After C-WAGS approval is received, return to Trial Details, publish the trial, and open entries when the host is ready.", "List Number"),
        (
            "Publishing is a separate step. Generating the application does not publish the trial, and the trial does not need to be published before the application is produced.",
            "Normal",
        ),
    )
    for text, style in application_content:
        publish_heading.insert_paragraph_before(text, style=style)
    application_picture = publish_heading.insert_paragraph_before()
    application_picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    application_picture.add_run().add_picture(
        str(annotated_by_title["Manual trial application reference"]),
        width=Inches(5.9),
    )
    application_caption = publish_heading.insert_paragraph_before(
        "Generate and submit the Trial Application while the trial remains in Draft; publish only after approval."
    )
    application_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    collaborator_content = (
        ("5.8 Trial Collaborators and Staff Roles", "Heading 2"),
        (
            "Trial Collaborators allow additional registered users to help with one trial without giving them access to every trial in the system. The person who created the trial is its owner; the owner or an application administrator controls collaborator invitations, role changes, and removal.",
            "Normal",
        ),
        ("Open Trial Collaborators from the selected trial's sidebar menu.", "List Number"),
        ("Enter the registered user's email address, choose the least-privileged role that fits the assignment, and send the invitation.", "List Number"),
        ("Ask the invited person to sign in and accept the invitation before relying on the assignment.", "List Number"),
        ("Remove or reduce access when the assignment ends.", "List Number"),
        (
            "Secretary: may edit trial setup; manage entries, waitlists, running order, and scores; manage financial records; and generate reports and the Trial Application. A secretary invited to an existing trial cannot manage collaborators or delete the trial. Those actions remain with the trial owner or an administrator.",
            "Normal",
        ),
        (
            "Assistant: may view the trial and help with entries, waitlists, running order, and score entry. An assistant cannot change trial setup, manage financial records, generate reports or the Trial Application, manage collaborators, or delete the trial.",
            "Normal",
        ),
        (
            "Use Assistant for event-day help and Secretary only when the person must configure the event or work with financial and official reporting functions. Collaborators never receive access to unrelated trials through this assignment.",
            "Normal",
        ),
    )
    for text, style in collaborator_content:
        publish_heading.insert_paragraph_before(text, style=style)

    postponed_day_content = (
        ("5.9 Moving a Postponed Trial Day", "Heading 2"),
        (
            "Use Edit Days to move an existing trial day to a replacement date. This is not a click-and-drag action. Change the date on the existing day so its classes, rounds, judge assignments, and entries remain attached.",
            "Normal",
        ),
        ("Open the trial from the dashboard and select Edit Days.", "List Number"),
        (
            "If the replacement date is outside the current trial range, change Start Date or End Date under Edit Trial Date Range, then select Update Date Range.",
            "List Number",
        ),
        (
            "In the Selected Days panel, find the postponed day and select its Trial day date field.",
            "List Number",
        ),
        ("Choose the replacement date from the date picker or type the new date.", "List Number"),
        ("Select Save Changes at the bottom of the page.", "List Number"),
        (
            "Return to the trial and verify the new date, classes, rounds, judges, and entry list before reopening entries or continuing event-day work.",
            "List Number",
        ),
        (
            "Do not remove the old day and create a new one when the day already has setup or entries. Editing the existing Trial day date preserves the day record and its attached information. Each trial day must use a unique date.",
            "Normal",
        ),
    )
    for text, style in postponed_day_content:
        publish_heading.insert_paragraph_before(text, style=style)

    for table in list(document.tables):
        title = placeholder_title(table)
        if title is None:
            continue
        if title == "Manual trial application reference":
            table._element.getparent().remove(table._element)
            continue
        if title in REMOVE_PLACEHOLDERS:
            table._element.getparent().remove(table._element)
            continue
        image_path = annotated_by_title.get(title)
        if image_path is None:
            table._element.getparent().remove(table._element)
            continue
        cell = table.cell(0, 0)
        clear_cell(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(4.25))

    # Footer and metadata.
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "C-WAGS Trial Secretary Manual | Secretary Edition"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(110, 110, 110)

    document.core_properties.title = "C-WAGS Trial Secretary Manual"
    document.core_properties.subject = "Secretary operations and event-day reference"
    document.core_properties.author = "C-WAGS Trial Management"
    document.core_properties.keywords = "C-WAGS, trial secretary, entries, waitlist, running order, scoring"
    document.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build()
